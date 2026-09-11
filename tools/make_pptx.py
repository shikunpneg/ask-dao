# -*- coding: utf-8 -*-
"""make_pptx.py — 用 python-pptx 生成可编辑的 .pptx（与 HTML 版同一份 deck_spec）。

要点：16:9；宣纸底；中文标题用宋体（必须同时设 latin 与 ea 两个字体槽，
否则 PowerPoint 里中文会静默回退成别的字）；表格用「No Style, Table Grid」拿细线，
不沿用自带配色；强调词用石青/朱砂区分（可复核的数用朱砂）。
"""
from __future__ import annotations

import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

INK, PAPER, MINERAL, VERM, BRONZE, SUB = "16191d", "f6f2e9", "2f6b6b", "b03a2e", "9a7b3f", "4a4f57"
SERIF, SANS = "SimSun", "Microsoft YaHei"
LAT_SERIF, LAT_SANS = "Source Han Serif SC", "Segoe UI"
SW, SH = Inches(13.333), Inches(7.5)
HAIR = Emu(9525)          # ≈0.75pt 细线


def _segments(s: str):
    """把 **粗** 与 `code` 拆成 [(文本, 强调?)]，pptx 里也保留强调。"""
    out, pos = [], 0
    for m in re.finditer(r"\*\*([^*]+)\*\*|`([^`]+)`", s):
        if m.start() > pos:
            out.append((s[pos:m.start()], False))
        out.append((m.group(1) or m.group(2), True))
        pos = m.end()
    if pos < len(s):
        out.append((s[pos:], False))
    return out


def style(run, *, size, bold=False, color=INK, serif=False, em=False):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.color.rgb = RGBColor.from_string(VERM if em else color)
    f.name = LAT_SERIF if serif else LAT_SANS
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", SERIF if serif else SANS)


def textbox(slide, x, y, w, h):
    tf = slide.shapes.add_textbox(x, y, w, h).text_frame
    tf.word_wrap = True
    return tf


def para(tf, first, segments, *, size, color=INK, serif=False, bold=False, space_after=6,
         em_color=None):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.space_after = Pt(space_after)
    for text, is_em in segments:
        r = p.add_run()
        r.text = text
        style(r, size=size, bold=bold or is_em,
              color=(em_color or (MINERAL if is_em else color)), serif=serif)
    return p


def bg(slide):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string(PAPER)


def rule(slide, x, y, w, color="c9c2b4", height=HAIR):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, height)
    sh.fill.solid()
    sh.fill.fore_color.rgb = RGBColor.from_string(color)
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def add_table(slide, x, y, w, rows, head, *, fs=10.5, head_fs=10.0, row_h=Inches(0.38)):
    t = slide.shapes.add_table(len(rows) + 1, len(head), x, y, w,
                               row_h * (len(rows) + 1)).table
    t.first_row = False
    t.horz_banding = False
    tblPr = t._tbl.tblPr
    for el in tblPr.findall(qn("a:tableStyleId")):
        tblPr.remove(el)
    el = tblPr.makeelement(qn("a:tableStyleId"), {})
    el.text = "{5940675A-B579-460E-94D1-54222C63F5DA}"     # No Style, Table Grid
    tblPr.append(el)
    for c, h in enumerate(head):
        cell = t.cell(0, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor.from_string("f7f3ea")
        p = cell.text_frame.paragraphs[0]
        p.space_after = Pt(0)
        r = p.add_run(); r.text = str(h)
        style(r, size=head_fs, bold=True, color=BRONZE)
    for ri, row in enumerate(rows, start=1):
        t.rows[ri].height = row_h
        for ci, val in enumerate(row):
            cell = t.cell(ri, ci)
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string("fffdf8")
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.space_after = Pt(0)
            for text, is_em in _segments(str(val)):
                r = p.add_run(); r.text = text
                style(r, size=fs, bold=is_em, color=(VERM if is_em else INK))
    return t


def build_pptx(slides, dst: Path) -> Path:
    prs = Presentation()
    prs.slide_width, prs.slide_height = SW, SH
    blank = prs.slide_layouts[6]
    n = len(slides)

    for i, s in enumerate(slides, start=1):
        sl = prs.slides.add_slide(blank)
        bg(sl)
        kind = s["kind"]
        cover = kind in ("cover", "end")

        tf = textbox(sl, Inches(0.75), Inches(1.28 if cover else 0.6), Inches(11.8), Inches(0.34))
        para(tf, True, [(s.get("kicker", ""), False)], size=10.5, color=BRONZE, bold=True,
             space_after=0)

        tf = textbox(sl, Inches(0.75), Inches(1.75 if cover else 0.98), Inches(11.8), Inches(1.7))
        for j, line in enumerate(s["title"].split("\n")):
            para(tf, j == 0, [(line, False)], size=(36 if cover else 27), serif=True, bold=True,
                 space_after=2)

        y = Inches(3.35 if cover else 2.0)
        if s.get("lead"):
            tf = textbox(sl, Inches(0.75), y, Inches(11.5), Inches(1.0))
            para(tf, True, _segments(s["lead"]), size=12, color=SUB, space_after=0)
            y = y + (Inches(1.0) if cover else Inches(0.92))

        if kind == "bullets":
            tf = textbox(sl, Inches(0.75), y, Inches(11.6), Inches(3.8))
            first = True
            for k, v in s["bullets"]:
                para(tf, first, [(k, False)], size=12.5, serif=True, bold=True, space_after=1)
                para(tf, False, _segments(v), size=11.5, color=SUB, space_after=9)
                first = False
        elif kind == "table":
            t = s["table"]
            rows = [[str(c) for c in r] for r in t["rows"]]
            fs = 10.5 if len(rows) <= 5 else (9.5 if len(rows) <= 7 else 9)
            add_table(sl, Inches(0.75), y, Inches(11.8), rows, [str(h) for h in t["head"]],
                      fs=fs, head_fs=10, row_h=Inches(0.3 if len(rows) > 6 else 0.42))
        elif kind == "chain":
            xs, ys = [Inches(0.75), Inches(6.9)], [y, y + Inches(1.55)]
            for kk, (k, v) in enumerate(s["steps"]):
                bx, by = xs[kk % 2], ys[kk // 2]
                rule(sl, bx, by, Inches(5.65), color="b03a2e", height=Pt(1.25))
                tf = textbox(sl, bx, by + Inches(0.08), Inches(5.65), Inches(0.3))
                para(tf, True, [(k, False)], size=10, color=BRONZE, bold=True, space_after=0)
                tf = textbox(sl, bx, by + Inches(0.4), Inches(5.65), Inches(1.1))
                para(tf, True, _segments(v), size=11, color=SUB, space_after=0)
        elif kind == "end":
            tf = textbox(sl, Inches(0.75), y, Inches(11.6), Inches(2.6))
            for kk, (k, v) in enumerate(s["links"]):
                p = para(tf, kk == 0, [(k + "　", False), (v, True)], size=11.5, color=SUB,
                         space_after=6, em_color=MINERAL)
                if p.runs:
                    p.runs[0].hyperlink.address = v
            tf2 = textbox(sl, Inches(0.75), Inches(6.3), Inches(11.6), Inches(0.4))
            para(tf2, True, [(s.get("foot", ""), False)], size=12, color=MINERAL, serif=True,
                 space_after=0)

        if s.get("note"):
            rule(sl, Inches(0.75), Inches(6.36), Inches(11.8))
            tf = textbox(sl, Inches(0.75), Inches(6.45), Inches(11.4), Inches(0.8))
            para(tf, True, _segments(s["note"]), size=10, color="6b7078", space_after=0)

        tf = textbox(sl, Inches(11.4), Inches(6.95), Inches(1.15), Inches(0.3))
        p = para(tf, True, [(f"{i:02d} / {n:02d}", False)], size=9.5, color=BRONZE, space_after=0)
        p.alignment = PP_ALIGN.RIGHT

    dst.parent.mkdir(parents=True, exist_ok=True)
    prs.save(dst)
    print(f"  PPTX: {dst.name}  {dst.stat().st_size:,} 字节  {n} 页")
    return dst
